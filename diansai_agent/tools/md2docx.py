#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown -> Word(.docx) 转换器（面向中文阅读文档排版）

支持：# / ## / ### 标题、GFM 表格（单元格内 <br> 换行、- [ ] 复选框）、
无序/有序列表、引用块、围栏代码块、分隔线、**粗体**、*斜体*、`行内代码`、[链接](url)。

用法:
    python md2docx.py 输入.md [输出.docx]
"""
from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "\u5fae\u8f6f\u96c5\u9ed1"      # 微软雅黑
MONO_FONT = "Consolas"
DARK = RGBColor(0x1F, 0x38, 0x64)
MID = RGBColor(0x2E, 0x54, 0x96)
GREY = RGBColor(0x50, 0x50, 0x50)
LINK = RGBColor(0x05, 0x63, 0xC1)
CODE = RGBColor(0xC7, 0x25, 0x4E)
HDR_FILL = "DCE6F1"
QUOTE_FILL = "F4F4F4"

TOKEN_RE = re.compile(
    r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*(?=[^\s*])[^*]*[^\s*]\*)"
)


# --------------------------------------------------------------------------- #
# 底层工具
# --------------------------------------------------------------------------- #
def set_run(run, size=10.5, bold=None, italic=None, color=None,
            font=BODY_FONT, mono=False):
    run.font.size = Pt(size)
    name = MONO_FONT if mono else font
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def shade(element, fill):
    pr = element.get_or_add_pPr() if element.tag.endswith("}p") else element
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def bottom_border(paragraph, size=8, color="1F3864", space=4):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


def left_border(paragraph, size=18, color="8EAADB"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)
    borders.append(left)
    ppr.append(borders)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_hyperlink(paragraph, url, text, size=10.5):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:rFonts", None),):
        pass
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), BODY_FONT)
    rfonts.set(qn("w:hAnsi"), BODY_FONT)
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run(run, size=9, color=GREY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# --------------------------------------------------------------------------- #
# 行内解析
# --------------------------------------------------------------------------- #
def add_inline(paragraph, text, size=10.5, bold=False, color=None):
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            set_run(paragraph.add_run(part[2:-2]), size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            set_run(paragraph.add_run(part[1:-1]), size=size - 0.5,
                    color=CODE, mono=True)
        elif part.startswith("[") and "](" in part:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if m:
                add_hyperlink(paragraph, m.group(2), m.group(1), size=size)
            else:
                set_run(paragraph.add_run(part), size=size, bold=bold, color=color)
        elif (part.startswith("*") and part.endswith("*") and len(part) > 2
              and not part.startswith("**")):
            set_run(paragraph.add_run(part[1:-1]), size=size, italic=True, color=color)
        else:
            set_run(paragraph.add_run(part), size=size, bold=bold, color=color)


def clean_cell_fragment(frag: str) -> str:
    frag = frag.strip()
    if frag.startswith("- [ ]"):
        return "\u2610 " + frag[5:].strip()
    if frag.lower().startswith("- [x]"):
        return "\u2611 " + frag[5:].strip()
    if frag.startswith("- "):
        return "\u2022 " + frag[2:].strip()
    return frag


# --------------------------------------------------------------------------- #
# 主转换
# --------------------------------------------------------------------------- #
def convert(md_path: str, docx_path: str) -> None:
    with open(md_path, "r", encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Cm(2.0))
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Cm(1.8))

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.25

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_p)

    i = 0
    n = len(lines)
    first_heading_done = False

    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        # 空行
        if not stripped:
            i += 1
            continue

        # 围栏代码块
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.0
            shade(p._p, "F5F5F5")
            for k, code_line in enumerate(buf):
                if k:
                    p.add_run().add_break()
                set_run(p.add_run(code_line), size=9, mono=True, color=RGBColor(0x33, 0x33, 0x33))
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            bottom_border(p, size=6, color="BFBFBF")
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not first_heading_done:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(8)
                add_inline(p, text, size=18)
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = DARK
                bottom_border(p, size=12, color="1F3864")
                first_heading_done = True
            elif level <= 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
                add_inline(p, text, size=14)
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = DARK
                bottom_border(p, size=6, color="8EAADB")
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                add_inline(p, text, size=12)
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = MID
            i += 1
            continue

        # 表格
        if stripped.startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for row_line in block:
                if re.fullmatch(r"\|[\s:\-|]+\|", row_line):
                    continue
                cells = [c for c in row_line.strip("|").split("|")]
                rows.append([c.strip() for c in cells])
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for ri, cells in enumerate(rows):
                row = table.add_row()
                for ci in range(cols):
                    cell = row.cells[ci]
                    cell.paragraphs[0].text = ""
                    frags = (cells[ci] if ci < len(cells) else "").split("<br>")
                    frags = [f for f in (clean_cell_fragment(x) for x in frags) if f]
                    if not frags:
                        frags = [""]
                    first = True
                    for frag in frags:
                        p = cell.paragraphs[0] if first else cell.add_paragraph()
                        first = False
                        p.paragraph_format.space_after = Pt(1)
                        p.paragraph_format.line_spacing = 1.1
                        add_inline(p, frag, size=9, bold=(ri == 0))
                    if ri == 0:
                        shade_cell(cell, HDR_FILL)
                if ri == 0:
                    repeat_header(row)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 引用块
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            for frag in buf:
                for sub in frag.split("<br>"):
                    if not sub.strip():
                        continue
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.45)
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    shade(p._p, QUOTE_FILL)
                    left_border(p)
                    add_inline(p, sub.strip(), size=10, color=GREY)
            continue

        # 复选框 / 无序列表
        m = re.match(r"^(\s*)[-*+]\s+(.*)$", raw)
        if m:
            indent = len(m.group(1))
            text = m.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5 + 0.5 * (indent // 2))
            p.paragraph_format.space_after = Pt(2)
            if text.startswith("[ ]"):
                add_inline(p, "\u2610 " + text[3:].strip(), size=10.5)
            elif text.lower().startswith("[x]"):
                add_inline(p, "\u2611 " + text[3:].strip(), size=10.5)
            else:
                add_inline(p, "\u2022 " + text, size=10.5)
            i += 1
            continue

        # 有序列表
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", raw)
        if m:
            indent = len(m.group(1))
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5 + 0.5 * (indent // 2))
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, f"{m.group(2)}. {m.group(3)}", size=10.5)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, stripped, size=10.5)
        i += 1

    doc.save(docx_path)


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: md2docx.py input.md [output.docx]")
        return 2
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(src)[0] + ".docx"
    convert(src, dst)
    print("OK ->", dst)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
