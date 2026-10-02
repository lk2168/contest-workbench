#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""表格列宽回归测试：中文被压成「一个字一行」这个 bug 不许再回来。

两个地方都要守住：
  ① **Word 文档**（md2docx）：列宽必须**写死**（tblLayout=fixed）、
     且每列不小于最小宽度、各列合计正好等于页面可用宽度
  ② **网页**（index.html）：表格要套 `.table-wrap`（可横向滚动）+ 单元格最小宽度

起因：用户贴了张截图 —— 赛题分析报告里「实际要求什么」那列被压成一个字一行。
根因是 `table.autofit = True`（Word/WPS 按内容自动分配，把短列压到极限）。
用法：python tests/test_report_table.py   （退出码 0 = 通过）
"""
from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from docx import Document                                              # noqa: E402
from docx.oxml.ns import qn                                            # noqa: E402
from docx.shared import Emu                                            # noqa: E402

from contest_workbench.tools.md2docx import (                          # noqa: E402
    算列宽, 视觉宽度, 表内字号, 设固定列宽, convert)

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(f"{'✅' if cond else '❌'} {name}" + (f"  —— {detail}" if detail and not cond else ""))


# 复刻用户截图里那张表：一列极长、一列极短，正是被压瘪的典型
超长表格 = """
| 来源 | 原文要点 | 实际要求什么 | 量化指标 | 备注 |
|---|---|---|---|---|
| 任务 | 载有平衡滚球装置的循线小车，沿黑环线行驶，球在装置内保持相对稳定 | 两套独立控制：①小车循迹定位 ②摆杆倾角到球位 | 原文未给 | 装置=小车+带凹槽摆杆+摆杆控制机构+钢球 |
| 任务 | 摆杆左端用铰链/合页固定，距小车平板高度 h≥5cm | 摆杆一端固定为转轴，另一端被电机顶起 | h≥5cm | 高度是硬指标 |
| 要求1 | 图传发送模块稳固装车，接收模块置环线外，实时显示并完整记录、可回放 | 一套无线图传+录像存储链路 | 画面要覆盖整个摆杆 | 6 分 |
"""


def main() -> int:
    print("== ① 视觉宽度（中文按两格算）==")
    check("全角汉字算 2 格", 视觉宽度("中文") == 4, str(视觉宽度("中文")))
    check("半角字母算 1 格", 视觉宽度("abc") == 3, str(视觉宽度("abc")))
    check("中英混排按实际算", 视觉宽度("h≥5cm") == 5, str(视觉宽度("h≥5cm")))
    check("Markdown 记号不占宽度", 视觉宽度("**粗体**") == 4, str(视觉宽度("**粗体**")))
    check("<br> 不占宽度", 视觉宽度("第一行<br>第二行") == 12, str(视觉宽度("第一行<br>第二行")))

    print("\n== ② 列宽分配（水位法）==")
    rows = [["任务", "很长的" * 30, "中等长度" * 5, "短", "备注" * 8],]
    宽 = 算列宽(rows, 5, 17.4, 9.0)
    check("★ 每列都不小于最小宽度 1.5cm", all(w >= 1.49 for w in 宽),
          str([round(w, 2) for w in 宽]))
    check("★ 各列合计正好等于可用宽度", abs(sum(宽) - 17.4) < 0.02,
          f"{sum(宽):.3f}")
    check("内容最长的那列拿到最多宽度", 宽[1] == max(宽), str([round(w, 2) for w in 宽]))
    check("比例正确：长列宽 ≈ 短列宽 × 内容比（180/40 = 4.5）",
          abs(宽[1] / 宽[2] - 4.5) < 0.3, f"{宽[1] / 宽[2]:.2f} vs 4.5")
    极端 = 算列宽([["a", "b"]], 2, 1.0, 9.0)         # 页面比两列最小宽度还窄
    check("页面极窄时平分成两列（不崩）", len(极端) == 2 and abs(sum(极端) - 1.0) < 1e-6,
          str(极端))
    check("列数为 0 时返回空", 算列宽([], 0, 17.4, 9.0) == [])

    print("\n== ③ 表内字号（内容太密自动缩）==")
    check("内容稀疏 → 用默认 9pt", 表内字号([["a", "b"]], 2, 17.4) == 9.0)
    密 = [[("很长的内容" * 40) for _ in range(5)] for _ in range(3)]
    check("★ 内容极密 → 自动缩到 7.5pt", 表内字号(密, 5, 17.4) == 7.5,
          str(表内字号(密, 5, 17.4)))
    check("列数为 0 时给默认值", 表内字号([], 0, 17.4) == 9.0)

    print("\n== ④ 真生成一份 docx 并读回验证（最关键）==")
    with tempfile.TemporaryDirectory() as td:
        md = Path(td) / "宽表.md"
        md.write_text("## 1. 任务与要求拆解\n" + 超长表格, encoding="utf-8")
        docx = Path(td) / "宽表.docx"
        convert(str(md), str(docx))
        check("docx 生成成功", docx.exists() and docx.stat().st_size > 0)

        doc = Document(str(docx))
        check("文档里有 1 张表", len(doc.tables) == 1, str(len(doc.tables)))
        tb = doc.tables[0]
        sec = doc.sections[0]
        可用 = Emu(sec.page_width - sec.left_margin - sec.right_margin).cm

        layout = tb._tbl.tblPr.find(qn("w:tblLayout"))
        typ = layout.get(qn("w:type")) if layout is not None else None
        check("★ 列宽是写死的（tblLayout=fixed），不再交给 Word 自动分配",
              typ == "fixed", str(typ))

        宽列表 = [c.width.cm if c.width is not None else None for c in tb.rows[0].cells]
        check("★ 每一列都有明确宽度", all(w is not None for w in 宽列表), str(宽列表))
        check("★ 没有列被压到 1.5cm 以下（就是「一个字一行」的判据）",
              all(w >= 1.49 for w in 宽列表), str([round(w, 2) for w in 宽列表]))
        check("★ 合计正好等于页面可用宽度（不留空、不溢出）",
              abs(sum(宽列表) - 可用) < 0.05, f"{sum(宽列表):.2f} vs {可用:.2f}")
        check("每一行的列宽都一致（不能只有表头行设了）",
              all(abs(r.cells[i].width.cm - 宽列表[i]) < 1e-6
                  for r in tb.rows for i in range(len(宽列表))), "")
        check("最长内容的列宽度最大（按内容分配，不是平均分）",
              宽列表[1] == max(宽列表), str([round(w, 2) for w in 宽列表]))
        # 页面可用宽度应约等于 17.4cm（A4 纵向 + 本项目页边距）
        check("页面可用宽度约 17.4cm（A4 纵向 + 1.8cm 边距）", abs(可用 - 17.4) < 0.1,
              f"{可用:.2f}")

    print("\n== ⑤ 网页那侧同样要守住 ==")
    html = (ROOT / "contest_workbench" / "web" / "static" / "index.html").read_text(
        encoding="utf-8")
    check("★ 表格外面套了 .table-wrap（宽度不够就整体横向滚动）",
          'class="table-wrap"' in html and "</table></div>" in html)
    check("★ CSS 里给了单元格最小宽度（中文不会被挤成竖排）",
          re.search(r"\.article td,\.article th\{min-width:\s*[\d.]+em", html) is not None)
    check("CSS 里有 .table-wrap 的横向滚动", ".table-wrap{overflow-x:auto" in html)
    check("不再用 word-break:break-all 处理表格（那会把中文按字断行）",
          "min-width:5.5em;word-break:normal" in html)

    print("\n" + "=" * 52)
    print(f"通过 {len(PASS)} 项，失败 {len(FAIL)} 项")
    if FAIL:
        print("失败项：" + "、".join(FAIL))
        return 1
    print("全部通过 ✅（Word 列宽 + 网页表格，都不许再压瘪中文）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
